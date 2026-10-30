#!/usr/bin/env python3
"""Faithfully sync a DOCX body to Markdown and verify coverage.

Usage:
  python docx_markdown_sync.py SOURCE.docx TARGET.md          # preview/check
  python docx_markdown_sync.py SOURCE.docx TARGET.md --write  # overwrite target
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph


def render_runs(paragraph: Paragraph) -> str:
    pieces: list[str] = []
    for run in paragraph.runs:
        text = run.text.replace("\n", "<br>")
        if not text:
            continue
        if run.bold and run.italic:
            text = f"***{text}***"
        elif run.bold:
            text = f"**{text}**"
        elif run.italic:
            text = f"*{text}*"
        pieces.append(text)
    return "".join(pieces).strip()


def escape_cell(text: str) -> str:
    return text.replace("\\", "\\\\").replace("|", "\\|").replace("\n", "<br>")


def render_cell(cell) -> str:
    parts = [render_runs(p) for p in cell.paragraphs if p.text.strip()]
    return escape_cell("<br>".join(parts))


def iter_blocks(document):
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def convert(document) -> str:
    lines: list[str] = []
    paragraph_index = 0
    for block in iter_blocks(document):
        if isinstance(block, Paragraph):
            raw = block.text.strip()
            if not raw:
                paragraph_index += 1
                continue
            text = render_runs(block) or raw
            style = block.style.name if block.style else ""
            if paragraph_index == 0:
                lines.extend([f"# {text}", ""])
            elif style.startswith("Heading "):
                level = int(style.split()[-1]) + 1
                lines.extend([f"{'#' * level} {text}", ""])
            elif style == "List Paragraph":
                lines.extend([f"- {text}", ""])
            else:
                lines.extend([text, ""])
            paragraph_index += 1
            continue

        rows = [[render_cell(cell) for cell in row.cells] for row in block.rows]
        if not rows:
            continue
        width = max(len(row) for row in rows)
        rows = [row + [""] * (width - len(row)) for row in rows]
        lines.append("| " + " | ".join(rows[0]) + " |")
        lines.append("| " + " | ".join(["---"] * width) + " |")
        lines.extend("| " + " | ".join(row) + " |" for row in rows[1:])
        lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def normalize(text: str) -> str:
    text = text.replace("\\|", "|").replace("<br>", " ")
    text = re.sub(r"[*_#|`]", "", text)
    text = re.sub(r"^\s*-\s+", "", text, flags=re.MULTILINE)
    return re.sub(r"\s+", " ", text).strip()


def verify(document, markdown: str) -> dict[str, int]:
    normalized = normalize(markdown)
    source_items = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            source_items.extend(cell.text for cell in row.cells if cell.text.strip())
    missing = [item for item in source_items if normalize(item) not in normalized]
    source_headings = sum(
        1 for p in document.paragraphs if p.style and p.style.name.startswith("Heading ")
    )
    markdown_headings = sum(
        1 for line in markdown.splitlines() if re.match(r"^#{2,10} ", line)
    )
    markdown_tables = sum(1 for line in markdown.splitlines() if line.startswith("| ---"))
    result = {
        "source_items": len(source_items),
        "missing_items": len(missing),
        "source_headings": source_headings,
        "markdown_headings": markdown_headings,
        "source_tables": len(document.tables),
        "markdown_tables": markdown_tables,
    }
    if missing:
        sample = "\n".join(f"- {item}" for item in missing[:10])
        raise ValueError(f"Markdown omitted {len(missing)} source items:\n{sample}")
    if source_headings != markdown_headings:
        raise ValueError(f"Heading-count mismatch: {result}")
    if len(document.tables) != markdown_tables:
        raise ValueError(f"Table-count mismatch: {result}")
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("target", type=Path)
    parser.add_argument("--write", action="store_true")
    args = parser.parse_args()

    document = Document(args.source)
    markdown = convert(document)
    result = verify(document, markdown)
    if args.write:
        args.target.parent.mkdir(parents=True, exist_ok=True)
        args.target.write_text(markdown, encoding="utf-8")
        result["bytes_written"] = len(markdown.encode("utf-8"))
    else:
        result["bytes_previewed"] = len(markdown.encode("utf-8"))
    print(result)


if __name__ == "__main__":
    main()
